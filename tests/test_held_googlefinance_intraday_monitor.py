# -*- coding: utf-8 -*-
"""
보유 8종목 GoogleFinance 장중감시 단위 테스트
- ReferencePrice(직전 장마감 종가) 및 ReferenceATR14 기반 ATRMove 검증
- anchor_price_p0(V4 앵커)와 장중 감시 기준값 분리 검증
- changepct / 100 서식 반영 및 거래정지(SUSPENDED) DataQuality 검증
"""

import unittest
import tempfile
from pathlib import Path
import openpyxl

from src.database.db_manager import DatabaseManager
from src.analysis.held_googlefinance_intraday_monitor import (
    get_held_8_universe_from_db,
    calc_atr_move,
    generate_held_8_intraday_excel,
    FULL_NAME_MAP,
    KOSDAQ_CODES
)

class TestHeldGoogleFinanceIntradayMonitor(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.db_path = str(Path(self.temp_dir.name) / "test_held.db")
        self.db = DatabaseManager(self.db_path)

        # 8개 보유종목 테스트 Fixture 등록
        # (code, name, qty, p0_anchor, a0_anchor, cur_atr14, last_close_ref_price, mode, mkt)
        sample_held = [
            ("000490", "대동", 4443, 7680.0, 645.2, 572.4, 7900.0, "NORMAL", "KOSPI"),
            ("004960", "한신공영", 2204, 11010.0, 578.9, 593.2, 12170.0, "NORMAL", "KOSPI"),
            ("010140", "삼성중공업", 495, 20150.0, 1052.1, 988.3, 21000.0, "NORMAL", "KOSPI"),
            ("055490", "테이팩스", 4283, 13760.0, 1074.2, 1060.2, 15840.0, "CONCENTRATION_RISK", "KOSPI"),
            ("161510", "PLUS 고배당주", 250, 24970.0, 705.6, 653.7, 25325.0, "NORMAL", "KOSPI"),
            ("206650", "유바이오로직스", 206, 8800.0, 639.5, 560.6, 8830.0, "NORMAL", "KOSDAQ"),
            ("234920", "자이글", 9314, 5310.0, 0.0, 0.0, 5310.0, "SUSPENDED_HOLD", "KOSDAQ"),
            ("490590", "RISE 미국AI밸류체인데일리고정커버드콜", 6271, 14330.0, 407.2, 390.1, 14385.0, "CONCENTRATION_RISK", "KOSPI"),
        ]

        for code, name, qty, p0, a0, cur_atr, ref_close, mode, mkt in sample_held:
            self.db.execute_non_query(
                "INSERT OR REPLACE INTO stock_info (stock_code, stock_name, market_type) VALUES (?, ?, ?)",
                (code, name, mkt)
            )
            self.db.execute_non_query(
                """
                INSERT OR REPLACE INTO portfolio_positions 
                (stock_code, quantity, avg_buy_price, anchor_price_p0, anchor_atr_a0, current_completed_atr, trade_mode)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (code, qty, p0, p0, a0, cur_atr, mode)
            )
            self.db.execute_non_query(
                """
                INSERT OR REPLACE INTO kiwoom_daily
                (stock_code, stk_date, open_price, high_price, low_price, close_price, volume)
                VALUES (?, '20260827', ?, ?, ?, ?, 100000)
                """,
                (code, ref_close, ref_close + 100, ref_close - 100, ref_close)
            )

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    def test_01_held_8_universe_extraction(self):
        """테스트 01: 실제 보유 8종목의 ReferencePrice(직전 종가) 및 ReferenceATR14 추출 검증"""
        universe = get_held_8_universe_from_db(db_manager=self.db)
        self.assertEqual(len(universe), 8)

        # 대동 검증: ReferencePrice는 직전 종가 7900원이며 P0 앵커 7680원과 명확히 구분됨
        daedong = next(u for u in universe if u["stock_code"] == "000490")
        self.assertEqual(daedong["ref_price"], 7900.0)
        self.assertEqual(daedong["p0_anchor"], 7680.0)
        self.assertEqual(daedong["ref_atr14"], 572.4)
        self.assertEqual(daedong["ticker"], "KRX:000490")

        # 한신공영 검증
        hanshin = next(u for u in universe if u["stock_code"] == "004960")
        self.assertEqual(hanshin["ref_price"], 12170.0)
        self.assertEqual(hanshin["p0_anchor"], 11010.0)
        self.assertEqual(hanshin["ref_atr14"], 593.2)

        for item in universe:
            self.assertEqual(len(item["stock_code"]), 6)
            self.assertGreater(item["quantity"], 0)
            self.assertIn(item["market"], ("KOSPI", "KOSDAQ"))
            self.assertIn("ref_price", item)
            self.assertIn("ref_atr14", item)
            self.assertIn("p0_anchor", item)

    def test_02_calc_atr_move_logic(self):
        """테스트 02: ReferencePrice 기준 ATRMove 및 가격상태(일반변동/의미있는 변화/강한 변화/결측) 판정 검증"""
        # 1) 일반변동 (< 0.5 ATR): 대동 직전 종가 7900원 기준 8050원 (+150원 -> +0.26 ATR)
        m_str, st, r = calc_atr_move(cur_price=8050, ref_price=7900.0, ref_atr14=572.4)
        self.assertIn("+0.26 ATR 상승", m_str)
        self.assertEqual(st, "일반변동")

        # 2) 의미있는 변화 (0.5 ~ 1.0 ATR): 대동 8300원 (+400원 -> +0.70 ATR)
        m_str, st, r = calc_atr_move(cur_price=8300, ref_price=7900.0, ref_atr14=572.4)
        self.assertIn("+0.70 ATR 상승", m_str)
        self.assertEqual(st, "의미있는 변화")

        # 3) 강한 변화 (>= 1.0 ATR): 한신공영 직전 종가 12170원 기준 11500원 (-670원 -> -1.13 ATR)
        m_str, st, r = calc_atr_move(cur_price=11500, ref_price=12170.0, ref_atr14=593.2)
        self.assertIn("-1.13 ATR 하락", m_str)
        self.assertEqual(st, "강한 변화")

        # 4) ATR 결측/0: 자이글
        m_str, st, r = calc_atr_move(cur_price=5310, ref_price=5310.0, ref_atr14=0.0)
        self.assertEqual(m_str, "ATR_DATA_MISSING")
        self.assertEqual(st, "ATR_DATA_MISSING")

    def test_03_generate_held_8_intraday_excel(self):
        """테스트 03: 보유 8종목 전용 엑셀(3개 탭) 생성 및 ReferencePrice / changepct / DataQuality 수식 검증"""
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out_file = Path(tmpdir) / "보유8종목_GoogleFinance_장중감시.xlsx"
            path = generate_held_8_intraday_excel(output_path=out_file, db_manager=self.db)
            self.assertTrue(path.exists())

            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["PORTFOLIO_CONFIG", "LIVE", "LIVE_SUMMARY"])

            # 1) PORTFOLIO_CONFIG check
            ws_cfg = wb["PORTFOLIO_CONFIG"]
            self.assertEqual(ws_cfg.max_row, 9)
            self.assertEqual(ws_cfg.cell(row=1, column=1).value, "Active")
            self.assertEqual(ws_cfg.cell(row=1, column=7).value, "ReferencePrice(직전종가)")
            self.assertEqual(ws_cfg.cell(row=1, column=8).value, "ReferenceATR14")
            self.assertEqual(ws_cfg.cell(row=1, column=9).value, "P0_Anchor(참고)")

            # 2) LIVE check
            ws_live = wb["LIVE"]
            self.assertEqual(ws_live.max_row, 9)
            self.assertIn("ReferencePrice", [ws_live.cell(row=1, column=c).value for c in range(1, 16)])
            self.assertIn("ReferenceATR14", [ws_live.cell(row=1, column=c).value for c in range(1, 16)])
            self.assertIn("ATRMove", [ws_live.cell(row=1, column=c).value for c in range(1, 16)])
            self.assertIn("가격상태", [ws_live.cell(row=1, column=c).value for c in range(1, 16)])
            self.assertIn('GOOGLEFINANCE(B2,"price")', str(ws_live.cell(row=2, column=3).value))
            # changepct / 100 check
            self.assertIn('GOOGLEFINANCE(B2,"changepct")/100', str(ws_live.cell(row=2, column=4).value))
            # volume ratio F>0 check
            self.assertIn('F2>0', str(ws_live.cell(row=2, column=11).value))
            # DataQuality check (SUSPENDED / DATA_REVIEW / VALID)
            self.assertIn("SUSPENDED", str(ws_live.cell(row=2, column=15).value))

            # 3) LIVE_SUMMARY check (9 columns)
            ws_sum = wb["LIVE_SUMMARY"]
            self.assertEqual(ws_sum.max_row, 9)
            expected_sum_headers = ["종목", "현재가", "등락률", "ATRMove", "가격상태", "거래량비", "당일고가", "당일저가", "DataQuality"]
            actual_sum_headers = [ws_sum.cell(row=1, column=c).value for c in range(1, 10)]
            self.assertEqual(actual_sum_headers, expected_sum_headers)
            self.assertEqual(ws_sum.cell(row=2, column=1).value, "=LIVE!A2")
            self.assertEqual(ws_sum.cell(row=2, column=4).value, "=LIVE!I2")

if __name__ == "__main__":
    unittest.main()

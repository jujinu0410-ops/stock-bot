# -*- coding: utf-8 -*-
"""
GoogleFinance 장중감시 Sheet Exporter (Sidecar) 단위 테스트
"""

import unittest
import tempfile
import os
from pathlib import Path
import openpyxl

from src.database.db_manager import DatabaseManager
from src.analysis.watchlist_googlefinance_exporter import (
    get_watchlist_universe_from_db,
    generate_googlefinance_watchlist_excel,
    export_watchlist_on_closing_report,
    FULL_NAME_MAP,
    KOSDAQ_CODES
)

class TestWatchlistGoogleFinanceExporter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.db_path = str(Path(self.temp_dir.name) / "test_watchlist.db")
        self.db = DatabaseManager(self.db_path)

        # Populate sample stocks in test stock_info
        sample_stocks = [
            ("005930", "삼성전자", "KOSPI"),
            ("140670", "알에스오토메이션", "KOSDAQ"),
            ("234920", "자이글", "KOSDAQ"),
            ("000490", "대동", "KOSPI"),
            ("004960", "한신공영", "KOSPI"),
        ]
        for code, name, mkt in sample_stocks:
            self.db.execute_non_query(
                "INSERT OR REPLACE INTO stock_info (stock_code, stock_name, market_type) VALUES (?, ?, ?)",
                (code, name, mkt)
            )

        self.sample_universe = [
            {"active": "Y", "stock_name": "삼성전자", "market": "KOSPI", "stock_code": "005930", "ticker": "KRX:005930"},
            {"active": "Y", "stock_name": "알에스오토메이션", "market": "KOSDAQ", "stock_code": "140670", "ticker": "KOSDAQ:140670"},
            {"active": "Y", "stock_name": "자이글", "market": "KOSDAQ", "stock_code": "234920", "ticker": "KOSDAQ:234920"},
        ]

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    def test_01_universe_structure_and_ticker_format(self):
        """테스트 01: 감시 종목 유니버스 추출 및 Ticker 포맷 검증"""
        universe = get_watchlist_universe_from_db(db_manager=self.db)
        self.assertIsInstance(universe, list)
        self.assertEqual(len(universe), 5)

        for item in universe:
            self.assertIn("active", item)
            self.assertIn("stock_name", item)
            self.assertIn("market", item)
            self.assertIn("stock_code", item)
            self.assertIn("ticker", item)

            code = item["stock_code"]
            self.assertEqual(len(code), 6)
            market = item["market"]
            self.assertIn(market, ("KOSPI", "KOSDAQ"))

            ticker = item["ticker"]
            if market == "KOSDAQ":
                self.assertTrue(ticker.startswith("KOSDAQ:"), f"KOSDAQ 티커 오류: {ticker}")
            else:
                self.assertTrue(ticker.startswith("KRX:"), f"KOSPI 티커 오류: {ticker}")

    def test_02_excel_generation_and_tabs(self):
        """테스트 02: 엑셀 파일 생성 및 WATCHLIST, LIVE 탭 구조 검증"""
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out_file = Path(tmpdir) / "보유관심종목_GoogleFinance_장중감시.xlsx"
            res_path = generate_googlefinance_watchlist_excel(
                output_path=out_file,
                universe=self.sample_universe
            )
            self.assertTrue(res_path.exists())

            wb = openpyxl.load_workbook(res_path)
            self.assertEqual(wb.sheetnames, ["WATCHLIST", "LIVE"])

            ws_w = wb["WATCHLIST"]
            self.assertEqual(ws_w.max_row, len(self.sample_universe) + 1)
            self.assertEqual(ws_w.cell(row=1, column=1).value, "Active")
            self.assertEqual(ws_w.cell(row=1, column=2).value, "종목명")
            self.assertEqual(ws_w.cell(row=1, column=3).value, "시장")
            self.assertEqual(ws_w.cell(row=1, column=4).value, "종목코드")
            self.assertEqual(ws_w.cell(row=1, column=5).value, "Ticker")

            # Sample row check
            self.assertEqual(ws_w.cell(row=2, column=1).value, "Y")
            self.assertEqual(ws_w.cell(row=2, column=2).value, "삼성전자")
            self.assertEqual(ws_w.cell(row=2, column=5).value, "KRX:005930")

    def test_03_live_tab_googlefinance_formulas(self):
        """테스트 03: LIVE 탭의 GOOGLEFINANCE 수식 정확성 검증"""
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out_file = Path(tmpdir) / "test_live_formulas.xlsx"
            generate_googlefinance_watchlist_excel(
                output_path=out_file,
                universe=self.sample_universe
            )
            wb = openpyxl.load_workbook(out_file)
            ws_l = wb["LIVE"]

            expected_headers = [
                "종목명", "Ticker", "현재가", "등락률", "거래량", "평균거래량", "거래량비",
                "당일고가", "당일저가", "52주고가", "52주고가거리", "TradeTime", "DataDelay", "DataQuality"
            ]
            actual_headers = [ws_l.cell(row=1, column=c).value for c in range(1, len(expected_headers) + 1)]
            self.assertEqual(actual_headers, expected_headers)

            # Row 2 formulas
            self.assertEqual(ws_l.cell(row=2, column=1).value, "=WATCHLIST!B2")
            self.assertEqual(ws_l.cell(row=2, column=2).value, "=WATCHLIST!E2")
            self.assertIn('GOOGLEFINANCE(B2,"price")', str(ws_l.cell(row=2, column=3).value))
            self.assertIn('GOOGLEFINANCE(B2,"changepct")/100', str(ws_l.cell(row=2, column=4).value))
            self.assertIn('GOOGLEFINANCE(B2,"volume")', str(ws_l.cell(row=2, column=5).value))
            self.assertIn('GOOGLEFINANCE(B2,"volumeavg")', str(ws_l.cell(row=2, column=6).value))
            self.assertIn('E2/F2', str(ws_l.cell(row=2, column=7).value))
            self.assertIn('F2>0', str(ws_l.cell(row=2, column=7).value))
            self.assertIn('GOOGLEFINANCE(B2,"high")', str(ws_l.cell(row=2, column=8).value))
            self.assertIn('GOOGLEFINANCE(B2,"low")', str(ws_l.cell(row=2, column=9).value))
            self.assertIn('GOOGLEFINANCE(B2,"high52")', str(ws_l.cell(row=2, column=10).value))
            self.assertEqual(ws_l.cell(row=2, column=11).value, '=IFERROR(C2/J2-1,NA())')
            self.assertIn('GOOGLEFINANCE(B2,"tradetime")', str(ws_l.cell(row=2, column=12).value))
            self.assertIn('GOOGLEFINANCE(B2,"datadelay")', str(ws_l.cell(row=2, column=13).value))
            self.assertIn("DATA_REVIEW", str(ws_l.cell(row=2, column=14).value))

    def test_04_closing_report_sidecar_invocation(self):
        """테스트 04: 장마감 사이드카 export_watchlist_on_closing_report 정상 호출 검증"""
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out_file = Path(tmpdir) / "sidecar_export.xlsx"
            path = export_watchlist_on_closing_report(db_manager=self.db, output_path=out_file)
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 1000)

if __name__ == "__main__":
    unittest.main()

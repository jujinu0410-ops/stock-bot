# -*- coding: utf-8 -*-
"""
보유 8종목 GoogleFinance 장중감시 모듈 (Sidecar)
- Source of Truth: portfolio_positions WHERE quantity > 0 (실제 보유 8종목 전용)
- 장중 감시용 기준값:
  * ReferencePrice: 직전 15:35 장마감 리포트의 종목별 마지막 가격 (kiwoom_daily 최신 종가)
  * ReferenceATR14: 같은 리포트에서 사용된 완성 ATR14 (portfolio_positions.current_completed_atr)
- ATRMove = (GoogleFinance 현재가 - ReferencePrice) / ReferenceATR14
- 등락률: GOOGLEFINANCE(..., "changepct")/100 서식 반영 (100배 중복 왜곡 방지)
- DataQuality: 거래정지(SUSPENDED), 결측(DATA_REVIEW), 정상(VALID) 3단계 분기
- 거래량비: 현재 거래량 / 평균 일거래량 (평균거래량 결측/0 시 NA() 처리)
- Shadow V1.0, ATR V4, 45M ADD ADVISORY 코어 100% 무영향 독립 사이드카
"""

from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.database.db_manager import DatabaseManager
from src.utils.logger import logger

# KOSDAQ 등록 종목 코드 세트
KOSDAQ_CODES = {
    "013030", "047770", "140670", "206650", "214450", "219550", "234920", "241520", "348340",
    "196170", "086520", "028300", "277810"
}

# 기본 종목명 보완 맵
FULL_NAME_MAP = {
    "000490": "대동",
    "004960": "한신공영",
    "005380": "현대차",
    "005930": "삼성전자",
    "009540": "HD한국조선해양",
    "010120": "LS일렉트릭",
    "010140": "삼성중공업",
    "011700": "한신기계",
    "012450": "한화에어로스페이스",
    "013030": "하이록코리아",
    "034020": "두산에너빌리티",
    "047050": "포스코인터내셔널",
    "047770": "코데즈컴바인",
    "055490": "테이팩스",
    "140670": "알에스오토메이션",
    "161510": "PLUS 고배당주",
    "206650": "유바이오로직스",
    "207940": "삼성바이오로직스",
    "214450": "파마리서치",
    "234920": "자이글",
    "241520": "DSC인베스트먼트",
    "267260": "HD현대일렉트릭",
    "348340": "뉴로메카",
    "371460": "TIGER 차이나전기차SOLACTIVE",
    "490590": "RISE 미국AI밸류체인데일리고정커버드콜",
}

def calc_atr_move(cur_price: Optional[float], ref_price: Optional[float], ref_atr14: Optional[float]) -> Tuple[str, str, Optional[float]]:
    """
    현재가, 직전 장마감 ReferencePrice, ReferenceATR14 기준 ATRMove 및 가격상태 판정
    - < 0.5: 일반변동
    - 0.5 ~ 1.0: 의미있는 변화
    - >= 1.0: 강한 변화
    - 결측/0: ATR_DATA_MISSING
    """
    if cur_price is None or ref_price is None or ref_atr14 is None:
        return "ATR_DATA_MISSING", "ATR_DATA_MISSING", None
    if ref_atr14 <= 0 or ref_price <= 0:
        return "ATR_DATA_MISSING", "ATR_DATA_MISSING", None

    diff = cur_price - ref_price
    move_ratio = diff / ref_atr14
    abs_move = abs(move_ratio)

    if move_ratio >= 0:
        move_str = f"+{move_ratio:.2f} ATR 상승"
    else:
        move_str = f"{move_ratio:.2f} ATR 하락"

    if abs_move < 0.5:
        state = "일반변동"
    elif abs_move < 1.0:
        state = "의미있는 변화"
    else:
        state = "강한 변화"

    return move_str, state, move_ratio

def get_held_8_universe_from_db(db_manager: Optional[DatabaseManager] = None) -> List[Dict[str, Any]]:
    """
    DB portfolio_positions(quantity > 0) 및 kiwoom_daily로부터 실제 보유 8종목의
    직전 장마감 ReferencePrice(종가) 및 ReferenceATR14 추출
    """
    db = db_manager if db_manager is not None else DatabaseManager()
    query = """
        SELECT 
            p.stock_code,
            s.stock_name,
            p.quantity,
            p.avg_buy_price,
            p.anchor_price_p0,
            p.anchor_atr_a0,
            p.current_completed_atr,
            p.confirmed_stop_price,
            p.trade_mode,
            s.market_type,
            (
                SELECT kd.close_price 
                FROM kiwoom_daily kd 
                WHERE kd.stock_code = p.stock_code 
                ORDER BY kd.stk_date DESC 
                LIMIT 1
            ) AS latest_close_price
        FROM portfolio_positions p
        LEFT JOIN stock_info s ON p.stock_code = s.stock_code
        WHERE p.quantity > 0
        ORDER BY p.stock_code
    """
    rows = db.execute_query(query)
    if not rows:
        return []

    held_universe = []
    for r in rows:
        r_dict = dict(r) if hasattr(r, "keys") else {
            "stock_code": r[0], "stock_name": r[1], "quantity": r[2], "avg_buy_price": r[3],
            "anchor_price_p0": r[4], "anchor_atr_a0": r[5], "current_completed_atr": r[6],
            "confirmed_stop_price": r[7], "trade_mode": r[8], "market_type": r[9],
            "latest_close_price": r[10] if len(r) > 10 else None
        }
        code = str(r_dict.get("stock_code", "")).strip().zfill(6)
        db_name = str(r_dict.get("stock_name", "")).strip()
        name = FULL_NAME_MAP.get(code, db_name if db_name and db_name != code else code)
        market = "KOSDAQ" if code in KOSDAQ_CODES else "KOSPI"
        ticker = f"KOSDAQ:{code}" if market == "KOSDAQ" else f"KRX:{code}"

        p0 = float(r_dict.get("anchor_price_p0") or 0.0)
        cur_atr = float(r_dict.get("current_completed_atr") or 0.0)
        a0 = float(r_dict.get("anchor_atr_a0") or 0.0)
        ref_atr14 = cur_atr if cur_atr > 0 else a0

        # ReferencePrice: 직전 장마감 종가 (없을 시 avg_buy_price / p0 대체)
        close_p = r_dict.get("latest_close_price")
        if close_p is not None and float(close_p) > 0:
            ref_price = float(close_p)
        elif p0 > 0:
            ref_price = p0
        else:
            ref_price = float(r_dict.get("avg_buy_price") or 0.0)

        held_universe.append({
            "active": "Y",
            "stock_name": name,
            "market": market,
            "stock_code": code,
            "ticker": ticker,
            "quantity": int(r_dict.get("quantity") or 0),
            "ref_price": ref_price,
            "ref_atr14": ref_atr14,
            "p0_anchor": p0,
            "trade_mode": str(r_dict.get("trade_mode", "NORMAL"))
        })

    return held_universe

def generate_held_8_intraday_excel(
    output_path: Optional[Path] = None,
    held_universe: Optional[List[Dict[str, Any]]] = None,
    db_manager: Optional[DatabaseManager] = None
) -> Path:
    """
    실제 보유 8종목 전용 GoogleFinance 장중감시 엑셀 템플릿 생성
    - Tab 1: PORTFOLIO_CONFIG (ReferencePrice 직전종가 및 ReferenceATR14)
    - Tab 2: LIVE (상세 실시간 감시표)
    - Tab 3: LIVE_SUMMARY (사용자 요청 9열 핵심 요약표)
    """
    if held_universe is None:
        held_universe = get_held_8_universe_from_db(db_manager)

    if output_path is None:
        output_path = Path("보유8종목_GoogleFinance_장중감시.xlsx")
    else:
        output_path = Path(output_path)

    wb = openpyxl.Workbook()

    # Sheet 1: PORTFOLIO_CONFIG
    ws_config = wb.active
    ws_config.title = "PORTFOLIO_CONFIG"

    # Sheet 2: LIVE
    ws_live = wb.create_sheet(title="LIVE")

    # Sheet 3: LIVE_SUMMARY
    ws_summary = wb.create_sheet(title="LIVE_SUMMARY")

    # Fonts & Fills
    font_header = Font(name="Malgun Gothic", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Malgun Gothic", size=10)
    font_bold = Font(name="Malgun Gothic", size=10, bold=True)
    fill_header_config = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    fill_header_live = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Slate
    fill_header_summary = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # 1. PORTFOLIO_CONFIG Tab
    config_headers = [
        "Active", "종목명", "시장", "종목코드", "Ticker", "보유수량",
        "ReferencePrice(직전종가)", "ReferenceATR14", "P0_Anchor(참고)", "매매모드"
    ]
    ws_config.append(config_headers)
    for col_idx in range(1, len(config_headers) + 1):
        cell = ws_config.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header_config
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_cell

    for row_idx, item in enumerate(held_universe, start=2):
        ws_config.append([
            item.get("active", "Y"),
            item.get("stock_name", ""),
            item.get("market", "KOSPI"),
            item.get("stock_code", ""),
            item.get("ticker", ""),
            item.get("quantity", 0),
            item.get("ref_price", 0.0),
            item.get("ref_atr14", 0.0),
            item.get("p0_anchor", 0.0),
            item.get("trade_mode", "NORMAL")
        ])
        fill_to_use = fill_zebra if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, len(config_headers) + 1):
            cell = ws_config.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
            cell.border = border_cell
            if col_idx in (1, 3, 4, 5, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            elif col_idx in (7, 8, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.0"

    # 2. LIVE Tab
    live_headers = [
        "종목", "Ticker", "현재가", "등락률", "거래량", "평균거래량",
        "ReferencePrice", "ReferenceATR14", "ATRMove", "가격상태", "거래량비", "당일고가", "당일저가", "TradeTime", "DataQuality"
    ]
    ws_live.append(live_headers)
    for col_idx in range(1, len(live_headers) + 1):
        cell = ws_live.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header_live
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_cell

    for i in range(2, len(held_universe) + 2):
        # A: 종목, B: Ticker, C: 현재가, D: 등락률(/100), E: 거래량, F: 평균거래량,
        # G: ReferencePrice, H: ReferenceATR14, I: ATRMove, J: 가격상태, K: 거래량비(E/F),
        # L: 당일고가, M: 당일저가, N: TradeTime, O: DataQuality
        formula_row = [
            f"=PORTFOLIO_CONFIG!B{i}",
            f"=PORTFOLIO_CONFIG!E{i}",
            f'=IFERROR(GOOGLEFINANCE(B{i},"price"),NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"changepct")/100,NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"volume"),NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"volumeavg"),NA())',
            f"=PORTFOLIO_CONFIG!G{i}",
            f"=PORTFOLIO_CONFIG!H{i}",
            f'=IF(OR(ISNA(C{i}), NOT(ISNUMBER(C{i})), G{i}<=0, H{i}<=0), "ATR_DATA_MISSING", IF((C{i}-G{i})/H{i}>=0, TEXT((C{i}-G{i})/H{i}, "+0.00") & " ATR 상승", TEXT((C{i}-G{i})/H{i}, "0.00") & " ATR 하락"))',
            f'=IF(OR(ISNA(C{i}), NOT(ISNUMBER(C{i})), G{i}<=0, H{i}<=0), "ATR_DATA_MISSING", IF(ABS((C{i}-G{i})/H{i})>=1.0, "강한 변화", IF(ABS((C{i}-G{i})/H{i})>=0.5, "의미있는 변화", "일반변동")))',
            f'=IFERROR(IF(AND(ISNUMBER(E{i}), ISNUMBER(F{i}), F{i}>0), E{i}/F{i}, NA()), NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"high"),NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"low"),NA())',
            f'=IFERROR(GOOGLEFINANCE(B{i},"tradetime"),NA())',
            f'=IF(PORTFOLIO_CONFIG!J{i}="SUSPENDED_HOLD", "SUSPENDED", IF(OR(NOT(ISNUMBER(C{i})), C{i}<=0, NOT(ISNUMBER(D{i})), NOT(ISNUMBER(E{i})), H{i}<=0), "DATA_REVIEW", "VALID"))'
        ]
        ws_live.append(formula_row)
        fill_to_use = fill_zebra if i % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, len(live_headers) + 1):
            cell = ws_live.cell(row=i, column=col_idx)
            cell.font = font_data
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
            cell.border = border_cell

            # Alignments
            if col_idx in (1,):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (2, 9, 10, 14, 15):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # Formats
            if col_idx in (3, 7, 8, 12, 13):
                cell.number_format = "#,##0"
            elif col_idx == 4:
                cell.number_format = "+0.00%;-0.00%;0.00%"
            elif col_idx in (5, 6):
                cell.number_format = "#,##0"
            elif col_idx == 11:
                cell.number_format = "0.00"

    # 3. LIVE_SUMMARY Tab (최종 9개 핵심 컬럼 표)
    # | 종목 | 현재가 | 등락률 | ATRMove | 가격상태 | 거래량비 | 당일고가 | 당일저가 | DataQuality |
    summary_headers = [
        "종목", "현재가", "등락률", "ATRMove", "가격상태", "거래량비", "당일고가", "당일저가", "DataQuality"
    ]
    ws_summary.append(summary_headers)
    for col_idx in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header_summary
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_cell

    for i in range(2, len(held_universe) + 2):
        summary_row = [
            f"=LIVE!A{i}",
            f"=LIVE!C{i}",
            f"=LIVE!D{i}",
            f"=LIVE!I{i}",
            f"=LIVE!J{i}",
            f"=LIVE!K{i}",
            f"=LIVE!L{i}",
            f"=LIVE!M{i}",
            f"=LIVE!O{i}"
        ]
        ws_summary.append(summary_row)
        fill_to_use = fill_zebra if i % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, len(summary_headers) + 1):
            cell = ws_summary.cell(row=i, column=col_idx)
            cell.font = font_data
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
            cell.border = border_cell

            # Alignments
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (4, 5, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # Formats
            if col_idx in (2, 7, 8):
                cell.number_format = "#,##0"
            elif col_idx == 3:
                cell.number_format = "+0.00%;-0.00%;0.00%"
            elif col_idx == 6:
                cell.number_format = "0.00"

    # Auto column width
    for ws in (ws_config, ws_live, ws_summary):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str.encode("utf-8", errors="ignore")))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"[HeldGoogleFinanceIntradayMonitor] 보유 8종목 장중감시 엑셀 템플릿 생성 완료: {output_path}")
    return output_path
